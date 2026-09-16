from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "知行数枢-企业数据智能运营中枢产品介绍.docx"
ARTIFACTS = ROOT / ".artifacts"
ARCH_IMAGE = ARTIFACTS / "product-architecture.png"

BLUE = "2E74B5"
DARK_BLUE = "17365D"
NAVY = "0B2545"
MID_BLUE = "4E86B4"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F4F8FC"
GOLD = "C6922C"
GRAY = "5B6573"
LIGHT_GRAY = "F2F4F7"
BORDER = "C9D3DF"
WHITE = "FFFFFF"
BLACK = "111827"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, color=BLACK, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def paragraph_border(paragraph, color=BORDER, size="8", side="left"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), "8")
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8.5, color=GRAY)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("知行数枢  |  企业数据智能运营中枢")
    set_run_font(r, size=8.5, bold=True, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_field(fp)


def add_numbering(doc: Document):
    numbering = doc.part.numbering_part.element

    def next_id(tag):
        values = [int(el.get(qn(f"w:{tag}Id"))) for el in numbering.findall(qn(f"w:{tag}"))]
        return max(values, default=0) + 1

    def make_num(fmt: str, text: str, font: str | None = None):
        abstract_id = next_id("abstractNum")
        num_id = next_id("num")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "279")
        p_pr.append(ind)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)
        return num_id

    return make_num("bullet", "•", "Symbol"), lambda: make_num("decimal", "%1.")


def add_list_item(doc, text: str, num_id: int):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths_dxa, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
        cell._tc.tcPr[-1].set(qn("w:fill"), LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=DARK_BLUE)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.05)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, PALE_BLUE)
    paragraph_border(p, BLUE, "16", "left")
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=11, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=BLACK)
    return p


def add_kicker(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    set_run_font(r, size=9.5, bold=True, color=GOLD)
    return p


def add_title(doc, text: str, size=30, color=NAVY, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_subtitle(doc, text: str, size=14, after=18):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size=size, color=GRAY)
    return p


def add_heading(doc, text: str, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True, color=DARK_BLUE)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def create_architecture_image(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1500, 980
    im = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(im)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    font = ImageFont.truetype(str(font_path), 34)
    small = ImageFont.truetype(str(font_path), 25)
    compact_small = ImageFont.truetype(str(font_path), 23)
    bold = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 36)
    title = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 42)
    draw.text((70, 38), "从企业系统到智能决策与执行的闭环", fill="#17365D", font=title)

    boxes = [
        ("企业业务系统", "吉客云 · CRM · 店铺 · 广告 · 财务 · 物流 · HR · 飞书", "#EAF2F8", "#2E74B5"),
        ("数据与内容接入", "API / 文件 / CDC / Webhook / 文档解析 / 同步任务", "#F2F4F7", "#5B6573"),
        ("企业数据智能中枢", "经营数据仓库 · 正式知识 · 角色记忆 · 主数据 · 指标语义", "#DDEBF7", "#17365D"),
        ("企业语义、证据与工具网关", "统一实体 / 统一指标 / 证据快照 / 内部API / MCP工具", "#E8F3EF", "#347A62"),
        ("智能体运行与应用", "管理者分身 · 数字会议 · 经营分析 · 运营助手 · 智能客服", "#FFF4DF", "#C6922C"),
        ("行动与复盘", "建议 → 人工确认 → 任务/操作 → 结果回写 → 经验沉淀", "#FCE9E6", "#B4574D"),
    ]
    x, box_w, box_h, gap = 120, 1120, 112, 24
    y = 118
    for i, (name, desc, fill, stroke) in enumerate(boxes):
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=20, fill=fill, outline=stroke, width=4)
        draw.text((x + 34, y + 20), name, fill=stroke, font=bold)
        desc_x = x + (555 if i == 3 else 350)
        desc_font = compact_small if i == 3 else small
        draw.text((desc_x, y + 31), desc, fill="#283444", font=desc_font)
        if i < len(boxes) - 1:
            arrow_x = x + box_w // 2
            draw.line((arrow_x, y + box_h, arrow_x, y + box_h + gap - 5), fill="#8795A6", width=5)
            draw.polygon([(arrow_x - 10, y + box_h + gap - 14), (arrow_x + 10, y + box_h + gap - 14), (arrow_x, y + box_h + gap)], fill="#8795A6")
        y += box_h + gap

    side_x = 1280
    draw.rounded_rectangle((side_x, 118, 1430, 766), radius=20, fill="#F7F9FB", outline="#9AA8B7", width=3)
    vertical_text = ["版本", "权限", "审计", "评测", "成本", "可观测"]
    yy = 160
    for t in vertical_text:
        draw.text((1320, yy), t, fill="#4B5968", font=font)
        yy += 96
    im.save(path, dpi=(180, 180))


def add_cover(doc: Document):
    for _ in range(3):
        doc.add_paragraph()
    add_kicker(doc, "企业私有化产品方案")
    add_title(doc, "知行数枢", size=34, after=2)
    add_title(doc, "企业数据智能运营中枢", size=25, color=BLUE, after=12)
    add_subtitle(doc, "把分散的数据、制度与组织经验，转化为可分析、可决策、可执行的企业智能能力。", size=14, after=28)

    add_table(
        doc,
        ["方案定位", "首期切入", "长期方向"],
        [["企业数据与智能体平台", "管理者智能分身", "企业Data & Agent OS"]],
        [3120, 3120, 3120],
        font_size=10,
    )
    add_callout(
        doc,
        "核心主张",
        "企业真正需要的不是更多聊天窗口，而是一套能持续读取最新企业事实、复用管理经验并推动业务闭环的数据智能中枢。",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(62)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("适用企业：电商、零售及多系统运营企业")
    set_run_font(r, size=10.5, bold=True, color=GRAY)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("产品方案 v0.1  |  2026年8月")
    set_run_font(r2, size=9.5, color=GRAY)
    doc.add_page_break()


def build_document():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    create_architecture_image(ARCH_IMAGE)
    doc = Document()
    configure_document(doc)
    bullet_id, new_decimal_id = add_numbering(doc)
    add_cover(doc)

    add_kicker(doc, "01  产品概述")
    add_title(doc, "让企业数据真正参与经营", size=24, after=10)
    add_body(doc, "企业已经拥有ERP、CRM、电商平台、广告平台、财务系统和大量文档，但这些资产仍然分散在不同系统和个人经验中。管理层看到的是彼此割裂的报表，员工遇到问题仍然反复询问，AI也无法稳定获取企业当前事实。")
    add_callout(doc, "知行数枢的答案", "先建立企业自己的数据、知识与证据底座，再让不同岗位的智能体基于同一事实工作，最终在人工确认下连接业务行动。")
    add_heading(doc, "客户将获得什么", 1)
    for item in [
        "统一企业经营事实：跨系统数据拥有统一实体、指标口径和历史版本。",
        "统一企业知识：制度、SOP和KPI按版本管理，回答可以定位到原文。",
        "复用组织经验：管理者和专家的判断方式沉淀为可审核的角色能力。",
        "提升决策效率：从找数、问人和开会汇总，升级为带证据的智能分析。",
        "形成执行闭环：AI先建议和草拟，成熟后在审批范围内执行重复工作。",
    ]:
        add_list_item(doc, item, bullet_id)
    add_heading(doc, "产品不是单点机器人", 1)
    add_body(doc, "管理者分身是第一个让企业快速感受到价值的应用，但平台的长期核心是企业数据智能中枢。相同的数据、知识、角色、工具和评测能力将继续支撑数字会议、经营分析、运营助手与智能客服。")
    doc.add_page_break()

    add_kicker(doc, "02  客户问题与产品回应")
    add_title(doc, "从数据孤岛到统一经营认知", size=24, after=12)
    add_table(
        doc,
        ["企业现状", "直接影响", "知行数枢的回应"],
        [
            ["数据分散", "报表口径冲突、取数慢", "统一实体、指标和证据快照"],
            ["制度频繁更新", "管理者重复解释", "版本化知识和带来源问答"],
            ["经验沉淀在个人", "人员变化后能力流失", "岗位资产与个人记忆分离"],
            ["会议缺少共同证据", "意见多、结论难复盘", "冻结数据版本的数字会议"],
            ["AI只会聊天", "无法进入真实业务流程", "MCP工具、行动中心与执行记录"],
        ],
        [2000, 3000, 4360],
    )
    add_heading(doc, "产品设计原则", 1)
    decimal_id = new_decimal_id()
    for idx, item in enumerate([
        "企业事实优先。当前制度和已确认经营数据高于历史聊天与模型常识。",
        "逻辑统一、存储分离。数据仓库、正式知识、角色记忆和运行记录各司其职。",
        "先给依据，再给判断。所有经营结论都说明指标口径、数据时间和来源。",
        "先建议，再执行。通过试点逐步提高自动化等级，不追求第一天无人值守。",
        "一个平台，多种角色。CEO、部门经理、员工和客服共享底座但获得不同体验。",
    ]):
        add_list_item(doc, item, decimal_id)
    doc.add_page_break()

    add_kicker(doc, "03  产品全景")
    add_title(doc, "六大中心，组成企业智能运营底座", size=24, after=10)
    add_table(
        doc,
        ["产品中心", "核心能力", "典型使用者"],
        [
            ["企业数据中心", "数据接入、实体、指标、质量与历史", "CEO、数据、各部门"],
            ["企业知识中心", "制度、SOP、KPI、FAQ和版本引用", "全体员工"],
            ["角色分身中心", "岗位模板、个人记忆、风格与委托边界", "管理者、专家"],
            ["数字会议中心", "独立分析、质询、决策包和行动项", "经营管理团队"],
            ["智能分析中心", "问数、预警、诊断、简报和运营指导", "CEO、经理、运营"],
            ["行动与执行中心", "工具、草拟、审批、执行和复盘", "运营、客服、管理层"],
        ],
        [1900, 5060, 2400],
    )
    add_heading(doc, "渠道入口", 1)
    add_body(doc, "员工可以在飞书和内部Web使用角色分身；管理者通过经营驾驶舱、数字会议和决策包工作；后续可接入企业微信、微信、店铺客服平台和业务软件。渠道只改变交互方式，不改变底层事实和工具。")
    add_callout(doc, "统一体验", "同一份制度、同一个指标、同一条角色记忆，在飞书、Web和数字会议中使用同一版本，避免不同AI给出不同企业答案。")
    doc.add_page_break()

    add_kicker(doc, "04  系统架构")
    add_title(doc, "数据中枢是核心，智能体是使用者", size=24, after=8)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(5)
    architecture_shape = p_img.add_run().add_picture(str(ARCH_IMAGE), width=Inches(6.35))
    architecture_shape._inline.docPr.set(
        "descr",
        "知行数枢总体逻辑架构：业务系统经过数据与内容接入进入企业数据智能中枢，"
        "再通过企业语义、证据与工具网关服务智能体应用，并形成行动与复盘闭环。",
    )
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    r = caption.add_run("图1  知行数枢总体逻辑架构")
    set_run_font(r, size=9, italic=True, color=GRAY)
    add_body(doc, "平台采用可插拔架构：知识、记忆、指标、智能体运行时和工作流都通过企业自有接口接入。企业可以先用轻量组件验证价值，再根据数据量和业务复杂度升级，不需要在首期一次性投入完整湖仓和大量微服务。")
    doc.add_page_break()

    add_kicker(doc, "05  首期切入")
    add_title(doc, "管理者智能分身：从重复答疑开始", size=24, after=10)
    add_body(doc, "管理者分身不是冒充本人发言，而是一个能够读取最新制度、经营数据和已审核管理经验的岗位代理。它首先处理高频重复问题，并把无法确认的事项交回本人。")
    add_heading(doc, "一个分身由四部分组成", 1)
    add_table(
        doc,
        ["组成", "内容"],
        [
            ["企业共同规则", "现行制度、KPI、术语和正式业务规则"],
            ["岗位资产", "职责、关注指标、决策边界和常用流程"],
            ["个人资产", "已审核偏好、历史案例、表达方式和反例"],
            ["当前上下文", "提问者、当前议题、实时数据和证据快照"],
        ],
        [2200, 7160],
    )
    add_heading(doc, "员工的一次提问", 1)
    decimal_id = new_decimal_id()
    for item in [
        "员工在飞书询问制度、KPI或经营问题。",
        "系统识别角色、问题类型和需要的资料。",
        "读取当前生效制度、指标和已审核角色记忆。",
        "分身给出事实、判断、未知项和引用证据。",
        "员工评价或纠错；重要事项可转交管理者确认。",
    ]:
        add_list_item(doc, item, decimal_id)
    add_callout(doc, "首期价值", "先减少管理者重复解释时间，同时验证知识版本、角色记忆、结构化数据、渠道接入和评测这五条长期主链路。")
    doc.add_page_break()

    add_kicker(doc, "06  数字会议")
    add_title(doc, "让分身围绕同一证据讨论", size=24, after=10)
    add_body(doc, "当平台拥有多个角色分身后，企业可以发起数字会议。CEO、运营、财务和反方智能体先独立分析，再进行交叉质询；主持智能体保留共识、分歧和未知项，最终由人类负责人确认。")
    add_heading(doc, "会议过程", 1)
    decimal_id = new_decimal_id()
    for item in [
        "登记议题、负责人、期限和成功指标。",
        "冻结制度、指标和经营数据版本。",
        "各角色独立分析，降低从众和迎合。",
        "两轮质询，所有主张绑定证据与假设。",
        "反方检查风险、错误激励和反事实。",
        "形成决策包，由人类确认并生成行动项。",
        "到期复盘，结果进入知识和记忆候选。",
    ]:
        add_list_item(doc, item, decimal_id)
    add_heading(doc, "首批会议议题", 1)
    add_table(
        doc,
        ["议题", "核心证据", "主要参与角色"],
        [
            ["是否增加广告预算", "ROI、毛利、库存、目标", "CEO / 运营 / 财务"],
            ["是否降价清理库存", "周转、库龄、毛利、现金", "CEO / 运营 / 供应链"],
            ["KPI是否造成错误激励", "制度、行为数据、部门反馈", "CEO / 人力 / 运营"],
        ],
        [2500, 4000, 2860],
    )
    doc.add_page_break()

    add_kicker(doc, "07  企业数据中心")
    add_title(doc, "不是复制所有系统，而是建立统一经营事实", size=24, after=10)
    add_body(doc, "知行数枢保留现有吉客云、CRM和店铺系统作为业务主系统。分析类数据同步到企业数据中心，强实时状态通过API查询，修改动作仍通过原系统正式接口完成，从而避免建设一个难维护的影子ERP。")
    add_heading(doc, "首个数据闭环", 1)
    add_table(
        doc,
        ["数据域", "首期指标", "可支持问题"],
        [
            ["销售", "GMV、订单量、客单价", "增长来自流量还是转化？"],
            ["退款", "退款额、退款率、原因", "异常来自商品还是服务？"],
            ["广告", "花费、成交、ROI", "预算是否应该增加？"],
            ["库存", "可售库存、库龄、周转", "哪些商品需要清理？"],
            ["利润", "毛利、贡献利润", "增长是否真正创造利润？"],
        ],
        [1500, 3460, 4400],
    )
    add_heading(doc, "未来扩展", 1)
    for item in [
        "继续接入CRM、客户、供应商、物流、财务和人力数据。",
        "形成企业主数据和指标目录，让所有报表与AI使用同一口径。",
        "通过MCP把经过治理的数据能力提供给Codex和其他智能体。",
        "为部门经理提供部门分析，为员工提供具体店铺运营指导。",
    ]:
        add_list_item(doc, item, bullet_id)
    doc.add_page_break()

    add_kicker(doc, "08  实施路线")
    add_title(doc, "全局架构，纵向切片", size=24, after=8)
    add_body(doc, "项目先完成统一边界、接口和管理壳子，再通过一个真实场景纵向贯通。这样既避免做成不可扩展的小机器人，也避免长期建设基础设施却无法证明业务价值。")
    add_table(
        doc,
        ["阶段", "核心交付", "客户可见结果"],
        [
            ["M0 平台骨架", "领域契约、管理台、数据与Agent接口", "看到完整产品结构和演进边界"],
            ["M1 分身试点", "制度、记忆、CEO分身、飞书、评测", "员工可自助问答，管理者可审核"],
            ["M2 数据与会议", "首个数据源、指标问数、三角色会议", "获得带数据证据的经营建议"],
            ["M3 运营执行", "每日巡店、行动中心、审批后操作", "AI发现问题并推动重复工作"],
            ["M4 客户服务", "客服工作台、订单物流工具、辅助回复", "提升客服效率并沉淀服务知识"],
        ],
        [1500, 4460, 3400],
    )
    add_heading(doc, "首期12周建议", 1)
    decimal_id = new_decimal_id()
    for item in [
        "第1-2周：确认总体架构、领域模型、技术栈和产品原型。",
        "第3-6周：完成知识、制度版本、记忆审核和角色分身。",
        "第7-9周：接入一个经营数据源，建立最小指标语义。",
        "第10-12周：飞书试点、真实问题评测、纠错和版本验收。",
    ]:
        add_list_item(doc, item, decimal_id)
    add_callout(doc, "交付原则", "每个阶段都必须有真实用户、真实数据和可量化验收，不以页面数量或模型演示作为完成标准。")
    doc.add_page_break()

    add_kicker(doc, "09  AI原生研发与开放架构")
    add_title(doc, "用Codex持续开发，也允许未来替换", size=24, after=10)
    add_body(doc, "项目采用Codex进行架构、编码、测试、文档和评测开发，并把重复研发流程固化为Skill。Codex Harness可作为首个智能体运行时，负责上下文、工具调用、进度事件和人工确认；业务领域通过自有接口与运行时解耦。")
    add_heading(doc, "可插拔能力", 1)
    add_table(
        doc,
        ["能力", "首期候选", "平台边界"],
        [
            ["知识", "WeKnora / RAGFlow", "通过KnowledgeProvider接入"],
            ["记忆", "自有审核流 / Tencent / Mem0", "通过MemoryProvider接入"],
            ["指标", "最小语义服务 / MetricFlow参考", "通过MetricService接入"],
            ["运行时", "Codex Runtime", "通过AgentRuntime接入"],
            ["工具", "企业MCP网关", "业务逻辑保留在内部服务"],
            ["工作流", "轻量队列 / Temporal候选", "按可靠性需求演进"],
        ],
        [1600, 3560, 4200],
    )
    add_body(doc, "这种方式让企业可以快速采用成熟开源能力，又不会把数据、角色、会议和业务动作绑定在单一项目上。")
    doc.add_page_break()

    add_kicker(doc, "10  试点与价值验证")
    add_title(doc, "先用一组可量化指标证明价值", size=24, after=10)
    add_heading(doc, "建议试点范围", 1)
    for item in [
        "一个公司或事业部。",
        "CEO、运营、财务三个角色。",
        "一个真实飞书试点群体。",
        "当前制度、KPI和20-50个高频问题。",
        "一个店铺或业务单元的日级经营数据。",
    ]:
        add_list_item(doc, item, bullet_id)
    add_heading(doc, "建议验收指标", 1)
    add_table(
        doc,
        ["指标", "试点目标", "说明"],
        [
            ["制度回答引用率", "100%", "每个制度结论都有版本和原文"],
            ["已知问题正确率", "≥95%", "按人工黄金集评测"],
            ["指标对账准确率", "≥99%", "与源系统或标准结果对账"],
            ["重复问询减少", "形成基线后目标≥50%", "统计转人工和管理者耗时"],
            ["员工反馈", "持续记录采纳与纠错", "不只统计点赞"],
        ],
        [2400, 2500, 4460],
    )
    add_callout(doc, "试点结论", "达到门槛后继续扩展数据源和角色；未达到时优先修正知识版本、指标口径和工作流，不用增加更多模型掩盖基础问题。")
    doc.add_page_break()

    add_kicker(doc, "11  未来蓝图")
    add_title(doc, "从企业信息中心走向智能运营系统", size=24, after=10)
    for item in [
        "员工：随时获得带依据的制度解释和岗位运营指导。",
        "部门经理：自动获取部门异常、经营原因和行动建议。",
        "CEO：基于统一数据获得决策分析、数字会议和复盘。",
        "运营团队：由AI完成每日巡店、报表汇总、任务草拟和低风险重复操作。",
        "客服团队：结合订单、商品和规则生成可靠回复，逐步自动处理低风险场景。",
        "企业组织：将数据、制度、经验和行动过程沉淀为持续进化的数字资产。",
    ]:
        add_list_item(doc, item, bullet_id)
    add_heading(doc, "最终形态", 1)
    add_body(doc, "知行数枢最终不是替代某一位员工或复制某一位领导，而是让整个企业拥有统一、持续更新、能够分析和推动工作的数字能力层。它连接已有系统，不推翻已有系统；复用人的经验，不取消人的责任；扩大AI能力，同时保留企业自己的数据和业务主导权。")
    add_callout(doc, "产品愿景", "数据归一，认知可信，决策有据，执行可控。")
    add_heading(doc, "建议下一步", 1)
    decimal_id = new_decimal_id()
    for item in [
        "确认首期试点部门、三类角色和高频问题清单。",
        "确认第一个结构化数据源及可用接入方式。",
        "完成M0架构决策与产品原型评审。",
        "启动12周管理者分身纵向切片。",
    ]:
        add_list_item(doc, item, decimal_id)

    core = doc.core_properties
    core.title = "知行数枢·企业数据智能运营中枢产品介绍"
    core.subject = "企业私有化数据、知识、智能体与运营执行平台"
    core.author = "知行数枢项目组"
    core.keywords = "企业数据中心, 管理者分身, 数字会议, 智能运营, Codex, MCP"
    core.comments = "客户产品介绍 v0.1"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
