import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

def create_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "业务确认"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY = "1F4E78"
    SECTION = "D9EAF7"
    NOTE_BG = "F3F6FA"
    WHITE = "FFFFFF"
    BAND = "F2F7FA"
    BORDER_COLOR = "B7C9D6"
    TEXT_COLOR = "203040"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    columns = [
        ("A", 16),
        ("B", 30),
        ("C", 44),
        ("D", 46),
        ("E", 40),
        ("F", 36),
        ("G", 20),
        ("H", 16),
        ("I", 40)
    ]
    for col_letter, width in columns:
        ws.column_dimensions[col_letter].width = width

    # Row 1: Title
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = "星云铁皮柜 Amazon 数据接入 — 核心业务确认表"
    title_cell.font = Font(name="微软雅黑", size=14, bold=True, color=WHITE)
    title_cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    # Row 2: Note 1
    ws.merge_cells("A2:I2")
    n1 = ws["A2"]
    n1.value = "说明：本表仅收录必须由公司业务、财务负责人决策的 4 项核心业务口径。接口传参、上下游单号关联及技术细节均已由系统自动闭环解决。"
    n1.font = Font(name="微软雅黑", size=10, color=TEXT_COLOR)
    n1.fill = PatternFill(start_color=NOTE_BG, end_color=NOTE_BG, fill_type="solid")
    n1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 25

    # Row 3: Note 2
    ws.merge_cells("A3:I3")
    n2 = ws["A3"]
    n2.value = "填写指引：请在“结论”列确认（如无异议选“确认”），并在“具体值/补充说明”列填写或调整具体值；若完全认可当前建议，直接回复确认即可。"
    n2.font = Font(name="微软雅黑", size=10, color=TEXT_COLOR)
    n2.fill = PatternFill(start_color=NOTE_BG, end_color=NOTE_BG, fill_type="solid")
    n2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 25

    # Row 4: Note 3 (Security)
    ws.merge_cells("A4:I4")
    n3 = ws["A4"]
    n3.value = "安全提示：严禁在本表中填写账号密码、AppSecret、Token、Cookie 等任何保密信息。系统已通过受控通道安全连接。"
    n3.font = Font(name="微软雅黑", size=9, color="7F8C8D")
    n3.fill = PatternFill(start_color=NOTE_BG, end_color=NOTE_BG, fill_type="solid")
    n3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[4].height = 22

    # Row 5: Empty spacing
    ws.row_dimensions[5].height = 10

    # Row 6: Table Headers
    headers = [
        "类别",
        "需要客户确认的事项",
        "目前系统掌握与默认情况",
        "请客户确认或填写",
        "填写示例（可直接参考）",
        "业务影响与用途",
        "建议由谁回复",
        "结论",
        "具体值/补充说明"
    ]
    ws.row_dimensions[6].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_idx, value=h)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Core 4 rows
    rows_data = [
        [
            "公司基本信息",
            "报表归属的公司全称与默认币种",
            "目前系统默认归入“星云项目主体（暂定）”，本位币为 USD，时区为北京时间（Asia/Shanghai）。",
            "请确认报表归属的法定主体全称、默认结算本位币（USD 或 CNY）以及所在时区。",
            "公司全称：XX科技有限公司；默认本位币：USD；业务时区：北京时间。",
            "决定经营驾驶舱与多中心报表的数据归属，以及金额与日期的统计基准。",
            "公司负责人 / 财务负责人",
            "确认",
            ""
        ],
        [
            "财务核算口径",
            "不同币种的金额换算规则",
            "当前北美 6 家店铺涉及 USD（美站）、CAD（加站）、MXN（墨站）。系统默认按财务每月月末汇率折算为本位币。",
            "请写明汇率来源渠道、按哪一天的汇率折算，多店铺间是否需要内部往来抵销，以及采用哪一版合并规则。",
            "汇率来源：公司财务统一提供（或中国银行折算汇率）；按月末汇率；无内部往来抵销；执行标准规则 V1。",
            "决定销售额、成本、毛利等关键指标在跨币种汇总时的财务准确性。",
            "财务负责人",
            "确认",
            ""
        ],
        [
            "历史数据起点",
            "历史经营数据回溯起始日期",
            "系统已调通订单、财务与广告历史链路；领星对不同模块有不同官方保留期，目前已验证 2025 年及 2026 年完整数据。",
            "请指定本次项目需要全量回溯的最早日期；超出领星平台保留范围的数据，系统将以平台最早可用日期为准。",
            "统一从 2024-01-01 开始回填历史；如果领星部分报表仅保留一年，则以领星最早可用日期为准。",
            "明确历史回填数据量与时间边界，避免将平台未保留的更早历史误判为数据同步遗漏。",
            "业务负责人 / 数据负责人",
            "确认",
            ""
        ],
        [
            "首次对账基准",
            "首次数据上线核对抽样数字",
            "系统已完成 6 家店铺的数据采集与规范事实映射，需要一个双方核准的基准完成端到端数据对账与验收。",
            "请提供同一店铺、同一日期范围（如已结账的 2026 年 8 月）由人工核对过的订单、销售额与退款汇总数字（可提供脱敏后台截图或报表导出）。",
            "2026年8月美站：总订单 1,234 单，销售额 USD 56,789，退款 23 单（附后台对账截图）。",
            "用于严格验证知行数枢抓取计算的指标与业务在领星后台看到的真实数字 100% 一致。",
            "业务负责人 / 财务负责人",
            "待提供",
            ""
        ]
    ]

    dv = DataValidation(type="list", formula1='"确认,待确认,调整,暂不适用"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H7:H10")

    for row_idx, row_values in enumerate(rows_data, 7):
        ws.row_dimensions[row_idx].height = 54
        fill_color = BAND if row_idx % 2 == 0 else WHITE
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="微软雅黑", size=10, color=TEXT_COLOR)
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx in (1, 7, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Freeze panes
    ws.freeze_panes = "A7"

    output_dir = r"D:\722\codex-tools\zhixing-data-hub\outputs\dat011"
    os.makedirs(output_dir, exist_ok=True)
    out_file1 = os.path.join(output_dir, "lingxing-amazon-业务数据确认表-精简交付版-20260916.xlsx")
    wb.save(out_file1)
    print(f"Successfully saved clean workbook to: {out_file1}")

    desktop_dir = r"C:\Users\41129\Desktop"
    if os.path.exists(desktop_dir):
        desktop_file = os.path.join(desktop_dir, "lingxing-amazon-业务数据确认表-精简交付版.xlsx")
        try:
            wb.save(desktop_file)
            print(f"Successfully saved clean workbook to Desktop: {desktop_file}")
        except Exception as e:
            print(f"Could not save to desktop: {e}")

if __name__ == "__main__":
    create_workbook()
